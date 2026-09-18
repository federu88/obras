# -*- coding: utf-8 -*-
"""Genera el SQL de carga del presupuesto de SR583 desde el Excel."""
import openpyxl, warnings, unicodedata, collections, io, datetime
warnings.filterwarnings("ignore")

BASE = r"C:\Users\feder\AppData\Local\Temp\claude\C--Users-feder-CLEAN-SEA-S-A-FILES-Terra-Mare---Documentos-Integra\0a9c950f-204b-4a63-87e6-d9dcf23d54a3\scratchpad"
P = BASE + r"\pres583.xlsx"
OUT = BASE + r"\carga-presupuesto-583.sql"

wb = openpyxl.load_workbook(P, data_only=True)
cat, dol = wb["Categoria"], wb["Dolar"]


def norm(s):
    if s is None:
        return ''
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    return ' '.join(s.lower().split())


def q(s):
    if s is None:
        return 'null'
    return "'" + str(s).strip().replace("'", "''") + "'"


# Nombres canonicos ya sembrados por 0006_seed_categories.sql
SEED_PADRES = ['Lote', 'Materiales Corralón/Hormigón', 'Sanitario', 'Electricidad',
               'Terminaciones', 'Mano de obra', 'Varios', 'Costos indirectos']
SEED_SUBS = ['Trámites', 'Pilotes', 'Columnas', 'Vigas + Losa S/PB',
             'Contrapisos y Carpetas', 'Mampostería', 'Revoques', 'Yesería + Pintura',
             'Losa Radiante', 'Instalación Agua / Sanitaria / Gas',
             'Instalación Eléctrica', 'Cables', 'Tablero', 'Pisos y Revestimientos',
             'Carpintería', 'Muebles', 'Herrería', 'Jornales', 'Contratistas', 'Otros',
             'Administración', 'Proyecto y arquitectura', 'Seguros', 'Gastos legales',
             'Comercialización', 'Financiación', 'Expensas', 'Gestoría']
MAP_P = {norm(x): x for x in SEED_PADRES}
MAP_S = {norm(x): x for x in SEED_SUBS}

# Tipeos del archivo que se corrigen al importar
FIX = {'mano de obra instalacion saniitaria': 'Mano de obra instalación sanitaria'}

UNID = {'gl': 'gl', 'uni': 'un', 'unidad': 'un', 'bol': 'bolsa',
        'rol': 'rollo', 'ml': 'ml', 'm2': 'm2'}
IGNORAR = {'', 'estimado', 'x', 'tbd', 'none'}


def canon_p(n):
    k = norm(n)
    if k in MAP_P:
        return MAP_P[k]
    n = str(n).strip()
    return n.title() if n.isupper() else n


def canon_s(n):
    k = norm(n)
    if k in FIX:
        return FIX[k]
    return MAP_S.get(k, str(n).strip())


# --- filas validas -----------------------------------------------------------
filas = []
for r in range(3, cat.max_row + 1):
    c, s, it = cat.cell(r, 1).value, cat.cell(r, 2).value, cat.cell(r, 3).value
    if not c or not it or norm(c) == 'x':
        continue
    filas.append(dict(
        r=r, cat=str(c).strip(), sub=str(s).strip() if s else 'Otros', item=str(it).strip(),
        prov=cat.cell(r, 4).value, uni=cat.cell(r, 5).value, mon=cat.cell(r, 6).value,
        st=cat.cell(r, 7).value, fec=cat.cell(r, 8).value,
        pu=cat.cell(r, 9).value, cant=cat.cell(r, 11).value))

cont = collections.Counter()
for f in filas:
    f['canon_cat'] = canon_p(f['cat'])
    f['canon_sub'] = canon_s(f['sub'])
    pre = ''.join(ch for ch in norm(f['canon_sub']) if ch.isalnum())[:4].upper() or 'ITEM'
    cont[pre] += 1
    f['code'] = "%s-%d" % (pre, cont[pre])
    f['unidad'] = UNID.get(norm(f['uni']), 'un')

provs = sorted({str(f['prov']).strip() for f in filas
                if f['prov'] and norm(f['prov']) not in IGNORAR})

fx = {}
for r in range(1, dol.max_row + 1):
    d, compra = dol.cell(r, 1).value, dol.cell(r, 2).value
    if isinstance(d, datetime.datetime) and isinstance(compra, (int, float)):
        fx.setdefault(d.date(), compra)

o = []
w = o.append
w("-- ===========================================================================")
w("-- Carga del presupuesto de SR583 desde 'Presupuesto 583_05.xlsx' (hoja Categoria)")
w("-- Generado automaticamente el %s. Correr UNA SOLA VEZ." % datetime.date.today())
w("--")
w("-- Los precios en pesos se convierten con public.fx_rate_at() sobre la fecha de")
w("-- cada fila, que arrastra la ultima cotizacion anterior o igual. Es lo que")
w("-- evita los 113 #N/A que tiene la planilla por 4 fechas faltantes.")
w("-- ===========================================================================")
w("")
w("begin;")
w("")
w("-- 1. Cotizaciones propias (columna Compra, el criterio de la planilla) -------")
w("insert into public.fx_rates (rate_date, source, ars_per_usd, note) values")
w(",\n".join("  ('%s', 'MEP', %s, 'Histórico planilla 583')" % (d, v)
             for d, v in sorted(fx.items())))
w("on conflict (rate_date, source) do nothing;")
w("")
w("-- 2. Categorias que falten ---------------------------------------------------")
for p in sorted({f['canon_cat'] for f in filas}):
    w("insert into public.cost_categories (name, kind)")
    w("  select %s, 'directo'" % q(p))
    w("  where not exists (select 1 from public.cost_categories"
      " where parent_id is null and name = %s);" % q(p))
for p, s in sorted({(f['canon_cat'], f['canon_sub']) for f in filas}):
    w("insert into public.cost_categories (parent_id, name, kind)")
    w("  select c.id, %s, 'directo' from public.cost_categories c" % q(s))
    w("  where c.parent_id is null and c.name = %s" % q(p))
    w("    and not exists (select 1 from public.cost_categories x"
      " where x.parent_id = c.id and x.name = %s);" % q(s))
w("")
w("-- 3. Proveedores -------------------------------------------------------------")
w("insert into public.suppliers (name) values")
w(",\n".join("  (%s)" % q(p) for p in provs))
w("on conflict (name) do nothing;")
w("")
w("-- 4. Items del catalogo ------------------------------------------------------")
for f in filas:
    w("insert into public.items (code, description, unit, category_id)")
    w("  select %s, %s, %s, s.id" % (q(f['code']), q(f['item']), q(f['unidad'])))
    w("  from public.cost_categories s"
      " join public.cost_categories p on p.id = s.parent_id")
    w("  where p.name = %s and s.name = %s" % (q(f['canon_cat']), q(f['canon_sub'])))
    w("  on conflict (code) do nothing;")
w("")
w("-- 5. Lineas de presupuesto de SR583 ------------------------------------------")
for f in filas:
    pu = f['pu'] if isinstance(f['pu'], (int, float)) else 0
    cant = f['cant'] if isinstance(f['cant'], (int, float)) else 1
    fecha = "'%s'" % f['fec'].date() if isinstance(f['fec'], datetime.datetime) else 'current_date'
    precio = ("%s" % pu) if f['mon'] == 'USD' else \
             "(%s / coalesce(public.fx_rate_at(%s), 1200))" % (pu, fecha)
    w("insert into public.budget_lines (project_id, item_id, category_id, description,"
      " unit, qty_original, price_original_usd)")
    w("  select pr.id, i.id, i.category_id, %s, %s, %s, round((%s)::numeric, 4)"
      % (q(f['item']), q(f['unidad']), cant, precio))
    w("  from public.projects pr, public.items i")
    w("  where pr.code = 'SR583' and i.code = %s;" % q(f['code']))
w("")
w("-- 6. Gastos reales: solo lo COMPRADO -----------------------------------------")
comprados = [f for f in filas if norm(f['st']) == 'comprado']
for f in comprados:
    pu = f['pu'] if isinstance(f['pu'], (int, float)) else 0
    cant = f['cant'] if isinstance(f['cant'], (int, float)) else 1
    fecha = "'%s'" % f['fec'].date() if isinstance(f['fec'], datetime.datetime) else 'current_date'
    mon = f['mon'] if f['mon'] in ('ARS', 'USD') else 'ARS'
    fxexpr = 'null' if mon == 'USD' else "coalesce(public.fx_rate_at(%s), 1200)" % fecha
    prov = q(f['prov']) if f['prov'] and norm(f['prov']) not in IGNORAR else 'null'
    w("insert into public.expenses (project_id, item_id, supplier_id, category_id,")
    w("       description, expense_date, qty, unit_price, currency, fx_usd, purchase_stage)")
    w("  select pr.id, i.id, sp.id, i.category_id, %s, %s, %s, %s, '%s', %s, 'recibida'"
      % (q(f['item']), fecha, cant, pu, mon, fxexpr))
    w("  from public.projects pr")
    w("  join public.items i on i.code = %s" % q(f['code']))
    w("  left join public.suppliers sp on sp.name = %s" % prov)
    w("  where pr.code = 'SR583';")
w("")
w("commit;")

io.open(OUT, 'w', encoding='utf-8').write("\n".join(o))

print("filas validas :", len(filas))
print("comprados     :", len(comprados))
print("proveedores   :", len(provs))
print("cotizaciones  :", len(fx))
print("categorias    :", len({f['canon_cat'] for f in filas}), "padres,",
      len({(f['canon_cat'], f['canon_sub']) for f in filas}), "subcategorias")
print("archivo       :", OUT)
