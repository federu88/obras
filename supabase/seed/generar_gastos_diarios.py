# -*- coding: utf-8 -*-
"""Genera el SQL de carga de los gastos diarios de los lotes 583 y 84."""
import openpyxl, warnings, unicodedata, datetime, io, re
warnings.filterwarnings("ignore")

B = r"C:\Users\feder\AppData\Local\Temp\claude\C--Users-feder-CLEAN-SEA-S-A-FILES-Terra-Mare---Documentos-Integra\0a9c950f-204b-4a63-87e6-d9dcf23d54a3\scratchpad"
OUT = B + r"\carga-gastos-diarios.sql"
INF = B + r"\gastos-sin-fecha.txt"

MESES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
         'julio': 7, 'agosto': 8, 'septiembre': 9, 'setiembre': 9, 'octubre': 10,
         'noviembre': 11, 'diciembre': 12}

# Seccion de la planilla -> (categoria padre, subcategoria)
SECCION = {
    'expensas':                    ('Costos indirectos', 'Expensas'),
    'gestoria':                    ('Costos indirectos', 'Gestoría'),
    'gestoria/extra':              ('Costos indirectos', 'Gestoría'),
    'gastos extras/ seguros':      ('Costos indirectos', 'Seguros'),
    'gastos edenor/naturgy':       ('Costos indirectos', 'Servicios de obra'),
    'gastos personal':             ('Mano de obra', 'Jornales'),
    'gastos obra':                 ('Varios', 'Gastos de obra'),
    'gastos tareas preliminares':  ('Varios', 'Tareas preliminares'),
}
NUEVAS = [('Costos indirectos', 'Servicios de obra', 'indirecto'),
          ('Varios', 'Gastos de obra', 'directo'),
          ('Varios', 'Tareas preliminares', 'directo')]

HOJAS = [('Gastos diarios lote 583', 'SR583'), ('Gastos diarios lote 84', 'SR084')]

# Ya importado desde el presupuesto: se saltea el duplicado exacto.
YA_CARGADO = {('gastos eidico', 93170.76)}


def norm(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode()
    return ' '.join(s.lower().split())


def q(s):
    return "'" + str(s).strip().replace("'", "''") + "'" if s is not None else 'null'


def mes_del_concepto(c):
    t = norm(c)
    for k, v in MESES.items():
        if k in t:
            return v
    return None


def resolver_fecha(fec, concepto):
    """Devuelve (fecha, nota) o (None, motivo)."""
    # Texto: quedo sin parsear porque el dia es > 12, asi que NO es ambiguo.
    if isinstance(fec, str):
        # Acepta d/m/aaaa y d-m-aa. Los años de dos digitos se completan a 20xx.
        m = re.match(r'^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2}|\d{4})\s*$', fec)
        if not m:
            return None, 'formato de texto no reconocido: %s' % fec
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return datetime.date(y, mo, d), 'texto d/m/a'
        except ValueError:
            return None, 'fecha invalida: %s' % fec

    if not isinstance(fec, datetime.datetime):
        return None, 'sin fecha'

    d = fec.date()
    mes = mes_del_concepto(concepto)
    # Invertida solo si el DIA guardado coincide con el mes nombrado y el mes no.
    if mes and d.day == mes and d.month != mes and d.month <= 12:
        try:
            return datetime.date(d.year, d.day, d.month), 'invertida'
        except ValueError:
            return d, 'ok'
    return d, 'ok'


wb = openpyxl.load_workbook(B + r"\gastos17.xlsx", data_only=True)

o, sin_fecha, invertidas, saltadas = [], [], 0, 0
w = o.append
w("-- ===========================================================================")
w("-- Gastos diarios de los lotes 583 y 84, desde 'Gastos SR84 y 583_Presupuesto1.7.xlsx'")
w("-- Generado el %s. Correr UNA SOLA VEZ." % datetime.date.today())
w("--")
w("-- Fechas: las que quedaron como TEXTO en el Excel se parsean dia/mes/año —")
w("-- quedaron sin parsear justamente porque el dia es mayor a 12, asi que no")
w("-- son ambiguas. Las que Excel si parseo se corrigen SOLO cuando el dia")
w("-- guardado coincide con el mes nombrado en el concepto y el mes no: dar")
w("-- vuelta todas romperia las 102 que estan bien.")
w("-- ===========================================================================")
w("")
w("begin;")
w("")
w("-- Categorias que faltan ------------------------------------------------------")
for padre, sub, kind in NUEVAS:
    w("insert into public.cost_categories (parent_id, name, kind)")
    w("  select c.id, %s, '%s' from public.cost_categories c" % (q(sub), kind))
    w("  where c.parent_id is null and c.name = %s" % q(padre))
    w("    and not exists (select 1 from public.cost_categories x"
      " where x.parent_id = c.id and x.name = %s);" % q(sub))
w("")

for hoja, code in HOJAS:
    ws = wb[hoja]
    w("-- %s -> %s %s" % (hoja, code, '-' * max(0, 40 - len(hoja))))
    seccion = None
    for r in range(1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        d = ws.cell(r, 4).value
        if b and isinstance(b, str):
            if norm(b).startswith('total'):
                continue
            seccion = b.strip()
        if not (isinstance(d, (int, float)) and d):
            continue
        if not c or norm(c).startswith('gastos -'):
            continue
        if (norm(c), round(d, 2)) in YA_CARGADO:
            saltadas += 1
            continue

        fecha, nota = resolver_fecha(ws.cell(r, 5).value, c)
        if fecha is None:
            sin_fecha.append("%s fila %s | %s | %s | %s" % (code, r, seccion, str(c)[:44], nota))
            continue
        if nota == 'invertida':
            invertidas += 1

        padre, sub = SECCION.get(norm(seccion), ('Varios', 'Otros'))
        w("insert into public.expenses (project_id, category_id, description,")
        w("       expense_date, qty, unit_price, currency, fx_usd, status)")
        w("  select pr.id, s.id, %s, '%s', 1, %s, 'ARS'," % (q(c), fecha, d))
        w("         coalesce(public.fx_rate_at('%s'), 1200), 'pagado'" % fecha)
        w("  from public.projects pr")
        w("  join public.cost_categories p on p.parent_id is null and p.name = %s" % q(padre))
        w("  join public.cost_categories s on s.parent_id = p.id and s.name = %s" % q(sub))
        w("  where pr.code = '%s';" % code)
    w("")

w("commit;")
io.open(OUT, 'w', encoding='utf-8').write("\n".join(o))
io.open(INF, 'w', encoding='utf-8').write("\n".join(sin_fecha) if sin_fecha else "(ninguna)")

inserts = sum(1 for l in o if l.startswith('insert into public.expenses'))
print("gastos a insertar      :", inserts)
print("fechas corregidas      :", invertidas)
print("saltadas por duplicado :", saltadas)
print("sin fecha utilizable   :", len(sin_fecha))
print("archivo                :", OUT)
